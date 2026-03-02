"""Generate formatted Excel workbooks from parsed MPesa statements."""

import io
import logging
from datetime import datetime
from decimal import Decimal
from typing import List, Dict

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

from app.services.statement_parser import (
    ParsedStatement,
    MPesaTransaction,
    TransactionType,
)

logger = logging.getLogger(__name__)


class MPesaExcelGenerator:
    """Generate professional Excel workbooks with charts."""

    COLORS = {
        "header_bg": "1B5E20",
        "header_text": "FFFFFF",
        "mpesa_green": "4CAF50",
        "alt_row": "F1F8E9",
        "positive": "E8F5E9",
        "negative": "FFEBEE",
        "total_row": "DCEDC8",
        "border": "BDBDBD",
    }

    def generate(self, statement: ParsedStatement) -> bytes:
        """Generate complete Excel file and return bytes."""
        wb = openpyxl.Workbook()

        self._create_summary_sheet(wb, statement)
        self._create_transactions_sheet(wb, statement)
        self._create_monthly_sheet(wb, statement)
        self._create_type_analysis_sheet(wb, statement)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        logger.info(f"Excel generated: {len(statement.transactions)} txns")
        return buffer.getvalue()

    # ─── Summary Sheet 

    def _create_summary_sheet(
        self, wb: openpyxl.Workbook, stmt: ParsedStatement
    ) -> None:
        ws = wb.active
        ws.title = "Summary"

        # Title
        ws.merge_cells("A1:E1")
        c = ws["A1"]
        c.value = "M-PESA STATEMENT ANALYSIS"
        c.font = Font(bold=True, size=18, color=self.COLORS["header_text"])
        c.fill = PatternFill("solid", fgColor=self.COLORS["header_bg"])
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        ws.merge_cells("A2:E2")
        ws["A2"].value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws["A2"].font = Font(italic=True, color="666666")
        ws["A2"].alignment = Alignment(horizontal="center")

        # Account Info
        row = 4
        self._section_header(ws, row, "A", "E", "ACCOUNT INFORMATION")
        row += 1

        info_items = [
            ("Account Name", stmt.metadata.account_name or "N/A"),
            ("Phone Number", stmt.metadata.phone_number or "N/A"),
            ("Total Transactions", str(stmt.metadata.transaction_count)),
        ]
        for label, val in info_items:
            ws[f"A{row}"] = label
            ws[f"A{row}"].font = Font(bold=True)
            ws[f"C{row}"] = val
            row += 1

        row += 1
        self._section_header(ws, row, "A", "E", "FINANCIAL SUMMARY")
        row += 1

        total_in = stmt.metadata.total_paid_in or Decimal("0")
        total_out = stmt.metadata.total_withdrawn or Decimal("0")
        net = total_in - total_out

        fin_items = [
            ("Total Money In (KES)", float(total_in)),
            ("Total Money Out (KES)", float(total_out)),
            ("Net Flow (KES)", float(net)),
        ]
        for label, val in fin_items:
            ws[f"A{row}"] = label
            ws[f"A{row}"].font = Font(bold=True)
            ws[f"C{row}"] = val
            ws[f"C{row}"].number_format = "#,##0.00"
            ws[f"C{row}"].font = Font(
                bold=True, color="2E7D32" if val >= 0 else "C62828"
            )
            row += 1

        # Type breakdown
        row += 1
        self._section_header(ws, row, "A", "E", "TRANSACTION TYPE BREAKDOWN")
        row += 1

        headers = ["Type", "Count", "Amount In", "Amount Out"]
        for col_i, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_i, value=h)
            cell.font = Font(bold=True, color=self.COLORS["header_text"])
            cell.fill = PatternFill("solid", fgColor=self.COLORS["mpesa_green"])
        row += 1

        type_totals = self._type_totals(stmt.transactions)
        for txn_type, data in sorted(
            type_totals.items(), key=lambda x: -x[1]["count"]
        ):
            ws.cell(row=row, column=1, value=txn_type.value)
            ws.cell(row=row, column=2, value=data["count"])
            c_in = ws.cell(row=row, column=3, value=float(data["paid_in"]))
            c_out = ws.cell(row=row, column=4, value=float(data["withdrawn"]))
            c_in.number_format = "#,##0.00"
            c_out.number_format = "#,##0.00"
            row += 1

        for col_letter, w in [("A", 30), ("B", 15), ("C", 20), ("D", 20), ("E", 15)]:
            ws.column_dimensions[col_letter].width = w

    # ─── Transactions Sheet ────────────────────────────────────

    def _create_transactions_sheet(
        self, wb: openpyxl.Workbook, stmt: ParsedStatement
    ) -> None:
        ws = wb.create_sheet("Transactions")

        headers = [
            "Receipt No.",
            "Date",
            "Time",
            "Description",
            "Type",
            "Counterparty",
            "Paid In (KES)",
            "Withdrawn (KES)",
            "Balance (KES)",
        ]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color=self.COLORS["header_text"])
            cell.fill = PatternFill("solid", fgColor=self.COLORS["header_bg"])
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        ws.row_dimensions[1].height = 30

        for row_idx, txn in enumerate(stmt.transactions, 2):
            # Determine row color
            if txn.paid_in and txn.paid_in > 0:
                bg = self.COLORS["positive"]
            elif txn.withdrawn and txn.withdrawn > 0:
                bg = self.COLORS["negative"]
            else:
                bg = self.COLORS["alt_row"] if row_idx % 2 == 0 else "FFFFFF"

            data = [
                txn.receipt_no,
                txn.completion_time.strftime("%Y-%m-%d") if txn.completion_time else "",
                txn.completion_time.strftime("%H:%M:%S") if txn.completion_time else "",
                txn.description[:120],
                txn.transaction_type.value,
                txn.counterparty[:50],
                float(txn.paid_in) if txn.paid_in else 0.0,
                float(txn.withdrawn) if txn.withdrawn else 0.0,
                float(txn.balance) if txn.balance else 0.0,
            ]

            for col, val in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.fill = PatternFill("solid", fgColor=bg)
                if col >= 7:
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                if col == 7 and isinstance(val, float) and val > 0:
                    cell.font = Font(color="1B5E20", bold=True)
                if col == 8 and isinstance(val, float) and val > 0:
                    cell.font = Font(color="B71C1C", bold=True)

        # Totals row
        total_row = len(stmt.transactions) + 2
        ws.cell(row=total_row, column=6, value="TOTALS").font = Font(bold=True)

        total_in = sum(float(t.paid_in) for t in stmt.transactions if t.paid_in)
        total_out = sum(float(t.withdrawn) for t in stmt.transactions if t.withdrawn)

        for col, val in [(7, total_in), (8, total_out)]:
            c = ws.cell(row=total_row, column=col, value=val)
            c.number_format = "#,##0.00"
            c.font = Font(bold=True, size=12)
            c.fill = PatternFill("solid", fgColor=self.COLORS["total_row"])

        widths = [18, 12, 10, 40, 18, 25, 16, 16, 16]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ─── Monthly Breakdown Sheet ───────────────────────────────

    def _create_monthly_sheet(
        self, wb: openpyxl.Workbook, stmt: ParsedStatement
    ) -> None:
        ws = wb.create_sheet("Monthly Breakdown")

        monthly: Dict[str, dict] = {}
        for txn in stmt.transactions:
            if not txn.completion_time:
                continue
            key = txn.completion_time.strftime("%Y-%m")
            if key not in monthly:
                monthly[key] = {
                    "label": txn.completion_time.strftime("%B %Y"),
                    "count": 0,
                    "paid_in": Decimal("0"),
                    "withdrawn": Decimal("0"),
                }
            monthly[key]["count"] += 1
            if txn.paid_in:
                monthly[key]["paid_in"] += txn.paid_in
            if txn.withdrawn:
                monthly[key]["withdrawn"] += txn.withdrawn

        headers = ["Month", "Transactions", "Paid In (KES)", "Withdrawn (KES)", "Net (KES)"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color=self.COLORS["header_text"])
            cell.fill = PatternFill("solid", fgColor=self.COLORS["header_bg"])

        for row_idx, (month_key, data) in enumerate(
            sorted(monthly.items()), 2
        ):
            net = data["paid_in"] - data["withdrawn"]
            row_data = [
                data["label"],
                data["count"],
                float(data["paid_in"]),
                float(data["withdrawn"]),
                float(net),
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                if col >= 3:
                    cell.number_format = "#,##0.00"
                if col == 5:
                    cell.font = Font(
                        bold=True,
                        color="1B5E20" if net >= 0 else "C62828",
                    )

        # Chart
        if len(monthly) > 1:
            chart = BarChart()
            chart.type = "col"
            chart.title = "Monthly Income vs Expense"
            chart.y_axis.title = "Amount (KES)"

            paid_ref = Reference(ws, min_col=3, max_col=3, min_row=1, max_row=len(monthly) + 1)
            withdrawn_ref = Reference(ws, min_col=4, max_col=4, min_row=1, max_row=len(monthly) + 1)
            months_ref = Reference(ws, min_col=1, min_row=2, max_row=len(monthly) + 1)

            chart.add_data(paid_ref, titles_from_data=True)
            chart.add_data(withdrawn_ref, titles_from_data=True)
            chart.set_categories(months_ref)
            chart.width = 20
            chart.height = 12

            ws.add_chart(chart, "G2")

        for col_letter, w in [("A", 20), ("B", 15), ("C", 18), ("D", 18), ("E", 15)]:
            ws.column_dimensions[col_letter].width = w

    # ─── Type Analysis Sheet ───────────────────────────────────

    def _create_type_analysis_sheet(
        self, wb: openpyxl.Workbook, stmt: ParsedStatement
    ) -> None:
        ws = wb.create_sheet("Type Analysis")

        headers = ["Type", "Count", "% Total", "Paid In", "Withdrawn", "Net"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color=self.COLORS["header_text"])
            cell.fill = PatternFill("solid", fgColor=self.COLORS["header_bg"])

        type_totals = self._type_totals(stmt.transactions)
        total_count = sum(d["count"] for d in type_totals.values())

        for row_idx, (txn_type, data) in enumerate(
            sorted(type_totals.items(), key=lambda x: -x[1]["count"]), 2
        ):
            pct = (data["count"] / total_count * 100) if total_count > 0 else 0
            net = data["paid_in"] - data["withdrawn"]

            row_data = [
                txn_type.value,
                data["count"],
                f"{pct:.1f}%",
                float(data["paid_in"]),
                float(data["withdrawn"]),
                float(net),
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                if col in [4, 5, 6]:
                    cell.number_format = "#,##0.00"

        for col_letter, w in [("A", 25), ("B", 10), ("C", 12), ("D", 18), ("E", 18), ("F", 15)]:
            ws.column_dimensions[col_letter].width = w

    # ─── Helpers ──────

    def _section_header(
        self, ws, row: int, start: str, end: str, title: str
    ) -> None:
        ws.merge_cells(f"{start}{row}:{end}{row}")
        cell = ws[f"{start}{row}"]
        cell.value = title
        cell.font = Font(bold=True, size=12, color=self.COLORS["header_text"])
        cell.fill = PatternFill("solid", fgColor=self.COLORS["mpesa_green"])
        ws.row_dimensions[row].height = 25

    def _type_totals(
        self, transactions: List[MPesaTransaction]
    ) -> Dict[TransactionType, dict]:
        totals: Dict[TransactionType, dict] = {}
        for txn in transactions:
            if txn.transaction_type not in totals:
                totals[txn.transaction_type] = {
                    "count": 0,
                    "paid_in": Decimal("0"),
                    "withdrawn": Decimal("0"),
                }
            totals[txn.transaction_type]["count"] += 1
            if txn.paid_in:
                totals[txn.transaction_type]["paid_in"] += txn.paid_in
            if txn.withdrawn:
                totals[txn.transaction_type]["withdrawn"] += txn.withdrawn
        return totals